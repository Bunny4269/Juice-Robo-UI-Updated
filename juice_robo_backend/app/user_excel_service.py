from pathlib import Path
from datetime import datetime, timezone

from openpyxl import Workbook, load_workbook


# ============================================================
# DEVELOPMENT USER CREDENTIALS EXCEL
# ============================================================

BASE_DIR = Path(__file__).resolve().parent.parent

DATA_DIR = BASE_DIR / "data"

EXCEL_FILE = DATA_DIR / "users.xlsx"

SHEET_NAME = "Users"

HEADERS = [
    "User ID",
    "Name",
    "Email",
    "Phone",
    "Password",
    "Status",
    "Is Admin",
    "Created At",
]


# ============================================================
# ENSURE EXCEL FILE EXISTS
# ============================================================

def ensure_user_excel_file():
    """
    Ensure the development user credentials Excel file exists.

    The file is created only once.

    If users.xlsx already exists, it is reused and nothing
    is overwritten.
    """

    # --------------------------------------------------------
    # Create data directory if required
    # --------------------------------------------------------

    DATA_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    # --------------------------------------------------------
    # Reuse existing Excel file
    # --------------------------------------------------------

    if EXCEL_FILE.exists():
        return EXCEL_FILE

    # --------------------------------------------------------
    # Create workbook
    # --------------------------------------------------------

    workbook = Workbook()

    worksheet = workbook.active

    worksheet.title = SHEET_NAME

    # --------------------------------------------------------
    # Header row
    # --------------------------------------------------------

    worksheet.append(
        HEADERS
    )

    # --------------------------------------------------------
    # Header formatting
    # --------------------------------------------------------

    for cell in worksheet[1]:
        cell.font = cell.font.copy(
            bold=True
        )

    worksheet.freeze_panes = "A2"

    # --------------------------------------------------------
    # Column widths
    # --------------------------------------------------------

    worksheet.column_dimensions["A"].width = 12
    worksheet.column_dimensions["B"].width = 25
    worksheet.column_dimensions["C"].width = 35
    worksheet.column_dimensions["D"].width = 18
    worksheet.column_dimensions["E"].width = 25
    worksheet.column_dimensions["F"].width = 15
    worksheet.column_dimensions["G"].width = 12
    worksheet.column_dimensions["H"].width = 28

    # --------------------------------------------------------
    # Save once
    # --------------------------------------------------------

    workbook.save(
        EXCEL_FILE
    )

    print(
        f"[Juice Robo] Created development "
        f"user Excel file: {EXCEL_FILE}"
    )

    return EXCEL_FILE


# ============================================================
# SAVE / UPDATE USER
# ============================================================

def save_user_credentials(
    *,
    user_id: int,
    name: str,
    email: str,
    phone: str,
    password: str,
    is_admin: bool = False,
    status: str = "ACTIVE",
    created_at=None,
):
    """
    Save or update a development user record.

    IMPORTANT:
    ------------------------------------------------------------
    The password stored in this Excel file is intentionally
    PLAINTEXT.

    This file is ONLY for local development/testing.

    The production/database password remains HASHED.
    ------------------------------------------------------------
    """

    # --------------------------------------------------------
    # Make sure the file exists
    # --------------------------------------------------------

    excel_path = ensure_user_excel_file()

    # --------------------------------------------------------
    # Timestamp fallback
    # --------------------------------------------------------

    if created_at is None:
        created_at = datetime.now(
            timezone.utc
        )

    # --------------------------------------------------------
    # Open existing workbook
    # --------------------------------------------------------

    workbook = load_workbook(
        excel_path
    )

    # --------------------------------------------------------
    # Get Users worksheet
    # --------------------------------------------------------

    if SHEET_NAME in workbook.sheetnames:

        worksheet = workbook[
            SHEET_NAME
        ]

    else:

        worksheet = workbook.create_sheet(
            SHEET_NAME
        )

        worksheet.append(
            HEADERS
        )

    # --------------------------------------------------------
    # Find existing user by User ID
    # --------------------------------------------------------

    existing_row = None

    for row in range(
        2,
        worksheet.max_row + 1,
    ):

        existing_id = worksheet.cell(
            row=row,
            column=1,
        ).value

        if (
            existing_id is not None
            and str(existing_id)
            == str(user_id)
        ):
            existing_row = row
            break

    # --------------------------------------------------------
    # Prepare row
    # --------------------------------------------------------

    values = [
        user_id,
        name,
        email,
        phone,

        # IMPORTANT:
        # Plain-text development password.
        password,

        status,
        bool(is_admin),

        (
            created_at.isoformat()
            if hasattr(
                created_at,
                "isoformat",
            )
            else str(created_at)
        ),
    ]

    # --------------------------------------------------------
    # Update existing user
    # --------------------------------------------------------

    if existing_row is not None:

        for column, value in enumerate(
            values,
            start=1,
        ):

            worksheet.cell(
                row=existing_row,
                column=column,
                value=value,
            )

    # --------------------------------------------------------
    # Add new user
    # --------------------------------------------------------

    else:

        worksheet.append(
            values
        )

    # --------------------------------------------------------
    # Save workbook
    # --------------------------------------------------------

    workbook.save(
        excel_path
    )

    return excel_path